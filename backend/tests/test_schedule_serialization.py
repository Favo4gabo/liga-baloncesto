from io import BytesIO

from openpyxl import Workbook

from types import SimpleNamespace

from app.services.schedule_extractor import (
    _filas_desde_ia,
    _mejor_coincidencia,
    _normalizar_medio,
    _catalogo_canonico,
    _validar_preview,
    _fecha_hora,
    _resolver_sugerencias,
    _normalizar_filas_excel,
    _fusionar_logistica_excel,
    _texto_completo_excel,
    _texto_csv_para_ia,
)


def test_excel_merged_values_are_expanded_and_rows_keep_coordinates(monkeypatch):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Jornada"
    hoja.merge_cells("A1:B1")
    hoja["A1"] = "Calendario"
    hoja.append(["Fecha", "Árbitro"])
    hoja.merge_cells("A3:A4")
    hoja["A3"] = "05/11/2026"
    hoja["B3"] = "Pedro"
    hoja["B4"] = "María"
    archivo = BytesIO()
    libro.save(archivo)

    monkeypatch.setenv("LLM_MAX_CHARS", "256")
    trozos = _texto_completo_excel(archivo.getvalue())
    texto = "\n".join(trozos)

    assert "R1: A=Calendario | B=Calendario" in texto
    assert "R3: A=05/11/2026 | B=Pedro" in texto
    assert "R4: A=05/11/2026 | B=María" in texto
    assert all(len(trozo) <= 256 for trozo in trozos)


def test_csv_chunks_keep_source_line_numbers_and_obey_limit(monkeypatch):
    monkeypatch.setenv("LLM_MAX_CHARS", "256")
    csv = ("FECHA,CANCHA,ARBITRO,NOTAS\n" + "05/11/2026,Forum,Pedro," + "dato " * 80 + "\n"
           + "06/11/2026,Domo,María,ok\n").encode()

    trozos = _texto_csv_para_ia(csv)
    texto = "\n".join(trozos)
    assert "R1: FECHA,CANCHA,ARBITRO,NOTAS" in texto
    assert "R3: 06/11/2026,Domo,María,ok" in texto
    assert all(len(trozo) <= 256 for trozo in trozos)


def test_ai_fields_and_transport_are_propagated():
    filas = _filas_desde_ia({"filas": [{
        "fecha": "05/11/2026", "hora": "18:30", "arbitro": "Pedro Martinez",
        "ciudad_origen": "Valencia", "ciudad_destino": "Caracas", "cancha": "Gimnasio",
        "medio_transporte": "terrestre", "cantidad_dias": 2, "cantidad_noches": 1,
        "viatico_por_dia": 25, "hospedaje_total": 55, "evidencia": "Jornada R8", "confianza": 0.91,
        "notas": "", "presupuesto": 170,
    }]})
    fila = filas[0]
    assert fila["ciudad_origen_texto"] == "Valencia"
    assert fila["ciudad_destino_texto"] == "Caracas"
    assert fila["medio_transporte"] == "taxi"
    assert fila["cantidad_dias"] == 2
    assert fila["hospedaje_total"] == 55
    assert fila["evidencia"] == "Jornada R8"
    assert _normalizar_medio("vuelo") == "avion"
    assert _normalizar_medio("desconocido") is None


def test_name_matching_supports_surname_first_and_rejects_ties():
    arbitros = [
        SimpleNamespace(id=1, nombre="María", apellido="González"),
        SimpleNamespace(id=2, nombre="Mario", apellido="González"),
    ]
    lookup = lambda item: f"{item.nombre} {item.apellido}"
    exact_id, score, candidates = _mejor_coincidencia("González, María", arbitros, lookup)
    assert exact_id == 1
    assert score == 1.0
    assert candidates[0]["id"] == 1

    tied_id, _, alternatives = _mejor_coincidencia("González", arbitros, lookup)
    assert tied_id is None
    assert len(alternatives) == 2


def test_canonical_prompt_catalog_is_capped_at_400_per_kind():
    ciudad = SimpleNamespace(id=1, nombre="Caracas", estado=SimpleNamespace(nombre="Distrito Capital"))
    arbitros = [SimpleNamespace(id=i, nombre=f"Nombre{i}", apellido="Árbitro", ciudad_origen=ciudad) for i in range(401)]
    canchas = [SimpleNamespace(id=i, nombre_cancha=f"Cancha{i}", ciudad=ciudad) for i in range(401)]
    ciudades = [SimpleNamespace(id=i, nombre=f"Ciudad{i}", estado=None) for i in range(401)]

    catalogo = _catalogo_canonico(arbitros, ciudades, canchas)
    assert all(len(items) == 400 for items in catalogo.values())
    assert catalogo["arbitros"][0]["ciudad_origen"] == "Caracas"
    assert catalogo["canchas"][0]["ciudad"] == "Caracas"


def test_preview_validation_marks_invalid_duplicate_and_low_confidence_rows():
    filas = [
        {"fecha": "2026-11-05", "hora": "18:30", "arbitro_id": 7, "confianza": 0.55},
        {"fecha": "2026-11-05", "hora": "18:30", "arbitro_id": 7, "confianza": 0.9},
        {"fecha": "2026-02-30", "hora": "25:61", "arbitro_texto": "Sin catálogo", "confianza": 0.9},
    ]

    validadas = _validar_preview(filas)
    assert validadas[0]["confianza_baja"] is True
    assert validadas[0]["marcada_revision"] is True
    assert validadas[0]["duplicado"] is True
    assert validadas[1]["duplicado"] is True
    assert validadas[2]["errores_validacion"] == [
        "fecha ausente o inválida; usa YYYY-MM-DD",
        "hora ausente o inválida; usa HH:MM en formato 24 horas",
    ]


def test_day_month_without_explicit_year_is_not_assigned_a_guessed_year():
    assert _fecha_hora("05/11", "18:30") == ("", "18:30")


def test_resolved_preview_exposes_requested_names_and_origin_city_id():
    ciudades = [SimpleNamespace(id=1, nombre="Valencia"), SimpleNamespace(id=2, nombre="Caracas")]
    arbitros = [SimpleNamespace(id=7, nombre="Pedro", apellido="Martínez", ciudad_origen_id=1)]
    canchas = [SimpleNamespace(id=9, nombre_cancha="Gimnasio", ciudad_id=2)]
    fila = _resolver_sugerencias({
        "fecha": "2026-11-05", "hora": "18:30", "arbitro_texto": "Martínez, Pedro",
        "ciudad_origen_texto": "", "ciudad_destino_texto": "Caracas", "ciudad_texto": "Caracas",
        "cancha_texto": "Gimnasio", "medio_transporte": "taxi", "evidencia": "Hoja1 R3",
    }, arbitros, canchas, ciudades)

    assert fila["arbitro"] == "Pedro Martínez"
    assert fila["ciudad_origen_id"] == 1
    assert fila["ciudad_origen"] == "Valencia"
    assert fila["ciudad_destino_id"] == 2
    assert fila["ciudad_destino"] == "Caracas"
    assert fila["cancha"] == "Gimnasio"


def test_explicit_lodging_cost_is_read_from_source_excel_and_overrides_ai_null():
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Untitled"
    hoja.append(["FECHA", "HORA", "ÁRBITRO", "CANCHA", "CIUDAD", "ORIGEN", "TIPO TRASLADO", "DÍAS VIÁTICO", "NOCHES", "HOSPEDAJE", "VIÁTICOS"])
    hoja.append(["2026-12-22", "19:30", "Andrés Silveira", "Gimnasio Ciudad de La Asunción", "Margarita", "Caracas", "Aéreo", 2, 2, 95, 65])
    archivo = BytesIO()
    libro.save(archivo)

    fuente = _normalizar_filas_excel(archivo.getvalue())
    previa_ia = [{
        "fecha": "2026-12-22", "hora": "19:30", "cancha_texto": "Gimnasio Ciudad de La Asunción",
        "evidencia": "Untitled · fila 2",
        "hospedaje_total": None, "cantidad_noches": None, "cantidad_dias": None,
        "viatico_por_dia": None, "medio_transporte": None,
    }]
    fusionadas = _fusionar_logistica_excel(previa_ia, fuente)[0]

    assert fuente[0]["hospedaje_total"] == 95
    assert fusionadas["hospedaje_total"] == 95
    assert fusionadas["cantidad_noches"] == 2
    assert fusionadas["cantidad_dias"] == 2
    assert fusionadas["viatico_por_dia"] == 32.5
    assert fusionadas["medio_transporte"] == "avion"
    assert fusionadas["ciudad_origen_texto"] == "Caracas"


def test_explicit_destination_city_does_not_inherit_wrong_city_from_existing_venue():
    ciudades = [SimpleNamespace(id=1, nombre="Barquisimeto"), SimpleNamespace(id=2, nombre="Barinas")]
    arbitros = [SimpleNamespace(id=7, nombre="Elena", apellido="Blanco", ciudad_origen_id=2)]
    canchas = [SimpleNamespace(id=9, nombre_cancha="Domo Bolivariano", ciudad_id=2)]
    fila = _resolver_sugerencias({
        "fecha": "2026-12-21", "hora": "18:00", "arbitro_texto": "Elena Blanco",
        "ciudad_texto": "Barquisimeto", "ciudad_destino_texto": "Barquisimeto",
        "ciudad_origen_texto": "Valencia", "cancha_texto": "Domo Bolivariano",
    }, arbitros, canchas, ciudades)

    assert fila["ciudad_id"] == 1
    assert fila["cancha_id"] is None
    assert fila["ciudad_destino"] == "Barquisimeto"
