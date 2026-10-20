"""Extracción de horarios Excel/CSV/imagen con sugerencias de catálogo."""

import io
import os
import re
import unicodedata
from datetime import date, datetime, time
from difflib import SequenceMatcher
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..models import Arbitro, Ciudad, UbicacionCancha
from .vision import extraer_filas_excel_texto, extraer_filas_imagen


ALIAS = {
    "fecha": ("fecha", "date", "dia", "día"),
    "hora": ("hora", "time", "horario"),
    "fecha_hora": ("fecha_hora", "fecha y hora", "fecha hora", "datetime", "inicio"),
    "arbitro": ("arbitro", "árbitro", "colegiado", "referee", "juez"),
    "cancha": ("cancha", "estadio", "sede", "venue", "pabellon"),
    "ciudad": ("ciudad", "localidad", "city", "municipio"),
    "ciudad_origen": ("ciudad_origen", "ciudad origen", "origen", "procedencia"),
    "tipo_traslado": ("tipo traslado", "medio transporte", "tipo de traslado"),
    "cantidad_dias": ("cantidad dias", "dias viatico", "dias de viatico"),
    "cantidad_noches": ("cantidad noches", "noches"),
    "hospedaje_total": ("hospedaje total", "hospedaje", "alojamiento"),
    "viaticos_total": ("viaticos total", "total viaticos", "viaticos"),
}


def _limpiar_texto(valor) -> str:
    if valor is None or pd.isna(valor):
        return ""
    return str(valor).strip()


def _texto_compacto(valor) -> str:
    """Normaliza saltos y espacios que suelen aparecer en celdas pegadas desde Excel."""
    return re.sub(r"\s+", " ", _limpiar_texto(valor))


def _clave(valor: str) -> str:
    texto = unicodedata.normalize("NFKD", str(valor).lower())
    texto = "".join(caracter for caracter in texto if not unicodedata.combining(caracter))
    return re.sub(r"[^a-z0-9]+", "", texto)


def _tokens(valor: str) -> set[str]:
    texto = unicodedata.normalize("NFKD", str(valor).lower())
    texto = "".join(caracter for caracter in texto if not unicodedata.combining(caracter))
    return set(re.findall(r"[a-z0-9]+", texto))


def _columna(columnas, tipo):
    alias = {_clave(item) for item in ALIAS[tipo]}
    for columna in columnas:
        if _clave(columna) in alias:
            return columna
    for columna in columnas:
        normalizada = _clave(columna)
        if any(item and (item in normalizada or normalizada in item) for item in alias if len(item) > 3):
            return columna
    return None


def _fecha_hora(fecha_valor, hora_valor="") -> tuple[str, str]:
    if isinstance(fecha_valor, (datetime, pd.Timestamp)):
        fecha_dt = pd.Timestamp(fecha_valor)
    elif isinstance(fecha_valor, date):
        fecha_dt = pd.Timestamp(fecha_valor)
    else:
        fecha_txt = _limpiar_texto(fecha_valor)
        if re.fullmatch(r"\d{1,2}[/-]\d{1,2}", fecha_txt):
            # Sin año solo la IA puede asociar uno si está explícito en otra parte del libro.
            return "", _normalizar_hora(hora_valor)
        formato_iso = bool(re.match(r"^\d{4}-\d{2}-\d{2}(?:[T ].*)?$", fecha_txt))
        fecha_dt = pd.to_datetime(fecha_txt, dayfirst=not formato_iso, errors="coerce") if fecha_txt else pd.NaT
    if pd.isna(fecha_dt):
        return "", _normalizar_hora(hora_valor)

    hora_txt = _normalizar_hora(hora_valor)
    if not hora_txt and isinstance(fecha_dt, pd.Timestamp) and (fecha_dt.hour or fecha_dt.minute):
        hora_txt = fecha_dt.strftime("%H:%M")
    return fecha_dt.strftime("%Y-%m-%d"), hora_txt


def _normalizar_hora(valor) -> str:
    if isinstance(valor, (datetime, pd.Timestamp, time)):
        return valor.strftime("%H:%M")
    texto = _limpiar_texto(valor)
    if not texto:
        return ""
    if isinstance(valor, (int, float)) and 0 <= valor < 1:
        minutos = round(float(valor) * 24 * 60)
        return f"{minutos // 60:02d}:{minutos % 60:02d}"
    try:
        convertido = pd.to_datetime(texto, errors="coerce")
        if not pd.isna(convertido):
            return convertido.strftime("%H:%M")
    except (TypeError, ValueError):
        pass
    match = re.search(r"\b(\d{1,2})[:h](\d{2})\b", texto.lower())
    if match and int(match.group(1)) < 24 and int(match.group(2)) < 60:
        return f"{int(match.group(1)):02d}:{int(match.group(2)):02d}"
    return ""


def _leer_tabular(contenido: bytes, extension: str):
    try:
        if extension == ".xlsx":
            return pd.read_excel(io.BytesIO(contenido), dtype=object).fillna("")
        return pd.read_csv(io.BytesIO(contenido), sep=None, engine="python", dtype=object).fillna("")
    except Exception as error:
        raise ValueError(f"No se pudo leer el archivo tabular: {error}") from error


def _normalizar_filas_excel(contenido: bytes) -> list[dict]:
    """Busca tablas de partidos en cualquier hoja y fila del libro.

    Muchos libros de liga agrupan varios catálogos en bloques de una sola hoja;
    por eso no se puede asumir que la fila 1 sea la cabecera del horario.
    """
    try:
        hojas = pd.read_excel(io.BytesIO(contenido), sheet_name=None, header=None, dtype=object)
    except Exception as error:
        raise ValueError(f"No se pudo leer el archivo Excel: {error}") from error

    filas = []
    for nombre_hoja, hoja in hojas.items():
        for indice_cabecera, serie in hoja.iterrows():
            valores = [_texto_compacto(valor) for valor in serie.tolist()]
            indices_columnas = [indice for indice, valor in enumerate(valores) if valor]
            columnas = [valores[indice] for indice in indices_columnas]
            mapa = {tipo: _columna(columnas, tipo) for tipo in ALIAS}
            if not (mapa["fecha"] or mapa["fecha_hora"]) or not mapa["cancha"]:
                continue

            registros = []
            for indice_fila in range(indice_cabecera + 1, len(hoja)):
                valores_fila = hoja.iloc[indice_fila].tolist()
                registro = {
                    columna: valores_fila[indice]
                    for indice, columna in zip(indices_columnas, columnas)
                    if indice < len(valores_fila)
                }
                # Los bloques del Excel de ejemplo están separados por una fila vacía.
                if not any(not _vacio_excel(valor) for valor in valores_fila):
                    if registros:
                        break
                    continue
                registros.append(registro)
            if registros:
                frame = pd.DataFrame(registros, columns=columnas)
                frame.index = range(indice_cabecera, indice_cabecera + len(frame))
                bloque = _normalizar_filas_tabulares(frame)
                for fila in bloque:
                    fila["hoja_origen"] = nombre_hoja
                    fila["evidencia"] = f"{nombre_hoja} · fila {fila['fila_origen']}"
                    filas.append(fila)
                # No volver a interpretar como encabezado el contenido de este bloque.
                break
    if not filas:
        raise ValueError(
            "No encontré una tabla de partidos en el Excel. Se requiere una fila de encabezados "
            "con fecha, cancha/estadio y, si están disponibles, hora, árbitro y ciudad."
        )
    return filas


def _vacio_excel(valor) -> bool:
    return valor is None or pd.isna(valor) or (isinstance(valor, str) and not valor.strip())


def _limite_llm() -> int:
    try:
        return max(256, int(os.getenv("LLM_MAX_CHARS", "12000")))
    except (TypeError, ValueError):
        return 12000


def _particionar_lineas(lineas: list[str], prefijo: str, limite: int) -> list[str]:
    """Agrupa líneas sin exceder el límite y preserva la cabecera en cada trozo."""
    trozos, actual = [], prefijo
    for linea in lineas:
        if len(actual) + len(linea) + 1 <= limite:
            actual += ("\n" if actual else "") + linea
            continue
        if actual.strip() and actual != prefijo:
            trozos.append(actual)
        cabecera = prefijo
        disponible = max(1, limite - len(cabecera) - 1)
        if len(linea) > disponible:
            partes = [linea[inicio:inicio + disponible] for inicio in range(0, len(linea), disponible)]
            for parte in partes[:-1]:
                trozos.append(f"{cabecera}\n{parte}"[:limite])
            actual = f"{cabecera}\n{partes[-1]}"[:limite]
        else:
            actual = f"{cabecera}\n{linea}" if cabecera else linea
    if actual.strip() and actual != prefijo:
        trozos.append(actual)
    return trozos or ([prefijo] if prefijo else [])


def _es_cabecera_tabular(linea: str) -> bool:
    texto = _clave(linea)
    tiene_fecha = any(alias in texto for alias in ("fecha", "date", "datetime", "dia"))
    tiene_partido = any(alias in texto for alias in ("cancha", "estadio", "sede", "venue", "arbitro", "referee", "colegiado"))
    return tiene_fecha and tiene_partido


def _prefijo_con_cabeceras(nombre: str, lineas: list[str], limite: int) -> str:
    prefijo = nombre
    max_contexto = max(0, limite // 3)
    for linea in lineas:
        if _es_cabecera_tabular(linea) and len(prefijo) + len(linea) + 12 <= max_contexto:
            prefijo += f"\nCABECERA: {linea}"
    return prefijo


def _texto_completo_excel(contenido: bytes) -> list[str]:
    """Serializa Excel con valores calculados, rangos combinados expandidos y chunks por hoja."""
    try:
        libro = load_workbook(BytesIO(contenido), data_only=True, read_only=False)
    except Exception as error:
        raise ValueError(f"No se pudo leer el archivo Excel: {error}") from error
    limite = _limite_llm()
    trozos = []
    for hoja in libro.worksheets:
        # En una combinación el valor solo vive en la celda superior izquierda.
        for rango in list(hoja.merged_cells.ranges):
            min_col, min_row, max_col, max_row = rango.bounds
            valor = hoja.cell(min_row, min_col).value
            hoja.unmerge_cells(str(rango))
            for fila in range(min_row, max_row + 1):
                for columna in range(min_col, max_col + 1):
                    hoja.cell(fila, columna).value = valor

        lineas = []
        for numero_fila in range(1, hoja.max_row + 1):
            celdas = []
            for numero_columna in range(1, hoja.max_column + 1):
                valor = hoja.cell(numero_fila, numero_columna).value
                if not _vacio_excel(valor):
                    letra = get_column_letter(numero_columna)
                    celdas.append(f"{letra}={_texto_compacto(valor)}")
            if celdas:
                lineas.append(f"R{numero_fila}: " + " | ".join(celdas))
        prefijo = _prefijo_con_cabeceras(f"HOJA: {hoja.title}", lineas, limite)
        trozos.extend(_particionar_lineas(lineas, prefijo, limite))
    libro.close()
    return trozos


def _texto_csv_para_ia(contenido: bytes) -> list[str]:
    texto = bytes(contenido).decode("utf-8-sig", errors="replace")
    lineas = [f"R{numero}: {linea}" for numero, linea in enumerate(texto.splitlines(), start=1) if linea.strip()]
    limite = _limite_llm()
    prefijo = _prefijo_con_cabeceras("CSV", lineas, limite)
    return _particionar_lineas(lineas, prefijo, limite)


def _catalogo_canonico(arbitros, ciudades, canchas) -> dict:
    """Forma el contexto de nombres canónicos, acotado para no desbordar el prompt."""
    return {
        "arbitros": [{
            "id": item.id,
            "nombre": item.nombre,
            "apellido": item.apellido,
            "ciudad_origen": getattr(getattr(item, "ciudad_origen", None), "nombre", None),
        } for item in arbitros[:400]],
        "ciudades": [{
            "id": item.id,
            "nombre": item.nombre,
            "estado": getattr(getattr(item, "estado", None), "nombre", None),
        } for item in ciudades[:400]],
        "canchas": [{
            "id": item.id,
            "nombre": item.nombre_cancha,
            "ciudad": getattr(getattr(item, "ciudad", None), "nombre", None),
        } for item in canchas[:400]],
    }


def _filas_desde_ia(resultado: dict) -> list[dict]:
    filas = []
    for indice, item in enumerate(resultado.get("filas", []), start=1):
        if not isinstance(item, dict):
            continue
        fecha, hora = _fecha_hora(item.get("fecha") or item.get("fecha_hora"), item.get("hora"))
        ciudad_destino = _limpiar_texto(item.get("ciudad_destino") or item.get("ciudad") or item.get("city"))
        evidencia = _limpiar_texto(item.get("evidencia"))
        arbitros = item.get("arbitro") or item.get("referee") or ""
        if isinstance(arbitros, (list, tuple)):
            nombres_arbitros = [_limpiar_texto(nombre) for nombre in arbitros if _limpiar_texto(nombre)]
        else:
            nombres_arbitros = [nombre.strip() for nombre in re.split(r"\s*(?:;|/|&|\by\b|\n)\s*", str(arbitros), flags=re.IGNORECASE) if nombre.strip()]
        if not nombres_arbitros:
            nombres_arbitros = [""]
        for nombre_arbitro in nombres_arbitros:
            fila = {
                "fila_origen": indice,
                "evidencia": evidencia or f"Fila de origen {indice}",
                "fecha": fecha,
                "hora": hora,
                "arbitro_texto": nombre_arbitro,
                "ciudad_origen_texto": _limpiar_texto(item.get("ciudad_origen")),
                "ciudad_destino_texto": ciudad_destino,
                "cancha_texto": _limpiar_texto(item.get("cancha") or item.get("estadio") or item.get("venue")),
                "ciudad_texto": ciudad_destino,
                "medio_transporte": _normalizar_medio(item.get("medio_transporte")),
                "_confianza_origen": item.get("confianza", 0.65),
            }
            hospedaje = item.get("hospedaje_total", item.get("hospedaje"))
            if hospedaje is not None:
                fila["hospedaje_total"] = hospedaje
            for campo in ("cantidad_dias", "cantidad_noches", "viatico_por_dia", "presupuesto", "notas"):
                if item.get(campo) is not None:
                    fila[campo] = item[campo]
            filas.append(fila)
    return filas


def _normalizar_medio(valor) -> str | None:
    texto = _clave(valor or "")
    if not texto:
        return None
    if any(alias in texto for alias in ("avion", "aereo", "vuelo", "air", "flight")):
        return "avion"
    if any(alias in texto for alias in (
        "terrestre", "taxi", "carretera", "bus", "autobus", "expreso",
        "vehiculoparticular", "carritoporpuesto", "transportelocal",
    )):
        return "taxi"
    return None


def _numero_opcional(valor):
    if valor is None or pd.isna(valor):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = re.sub(r"[^0-9,.-]", "", str(valor).strip())
    if not texto:
        return None
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".") if texto.rfind(",") > texto.rfind(".") else texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def _mejor_coincidencia(texto: str, opciones, etiqueta) -> tuple[int | None, float, list[dict]]:
    original = _texto_compacto(texto)
    consulta = _clave(original)
    if not consulta:
        return None, 0.0, []

    # Algunos libros ya contienen los IDs internos en vez de los nombres.
    # Excel suele cargarlos como número entero o decimal terminado en .0.
    if re.fullmatch(r"[1-9]\d*(?:\.0+)?", original):
        id_directo = int(float(original))
        if any(opcion.id == id_directo for opcion in opciones):
            opcion = next(opcion for opcion in opciones if opcion.id == id_directo)
            return id_directo, 1.0, [{"id": id_directo, "texto": etiqueta(opcion), "confianza": 1.0}]

    consulta_nombre = original
    if "," in original:
        apellido, nombre = [parte.strip() for parte in original.split(",", 1)]
        consulta_nombre = f"{nombre} {apellido}"
    consulta_clave = _clave(consulta_nombre)
    tokens_consulta = _tokens(consulta_nombre)
    puntuaciones = []
    for opcion in opciones:
        nombre_opcion = _texto_compacto(etiqueta(opcion))
        candidato = _clave(nombre_opcion)
        tokens_candidato = _tokens(nombre_opcion)
        score = SequenceMatcher(None, consulta_clave or consulta, candidato).ratio() if candidato else 0
        if consulta_clave == candidato or consulta == candidato:
            score = 1.0
        elif consulta_clave and (consulta_clave in candidato or candidato in consulta_clave):
            score = max(score, 0.82)
        if tokens_consulta and tokens_candidato:
            cobertura = len(tokens_consulta & tokens_candidato) / len(tokens_consulta | tokens_candidato)
            score = max(score, cobertura)
        puntuaciones.append((opcion.id, score, nombre_opcion))
    puntuaciones.sort(key=lambda candidato: (-candidato[1], candidato[2].casefold()))
    alternativas = [
        {"id": opcion_id, "texto": nombre, "confianza": round(score, 3)}
        for opcion_id, score, nombre in puntuaciones[:3]
        if score >= 0.30
    ]
    if not puntuaciones or puntuaciones[0][1] < 0.55:
        return None, round(puntuaciones[0][1], 3) if puntuaciones else 0.0, alternativas
    mejor_id, mejor_score, _ = puntuaciones[0]
    # Si dos opciones son prácticamente igual de plausibles, exige selección humana.
    if len(puntuaciones) > 1 and mejor_score - puntuaciones[1][1] < 0.035:
        return None, round(mejor_score, 3), alternativas
    return mejor_id, round(mejor_score, 3), alternativas


def _resolver_sugerencias(fila: dict, arbitros, canchas, ciudades) -> dict:
    arbitro_id, score_arbitro, candidatos_arbitro = _mejor_coincidencia(
        fila["arbitro_texto"], arbitros, lambda item: f"{item.nombre} {item.apellido}"
    )
    ciudad_texto = fila.get("ciudad_destino_texto") or fila.get("ciudad_texto", "")
    ciudad_id, score_ciudad, candidatos_ciudad = _mejor_coincidencia(
        ciudad_texto, ciudades, lambda item: item.nombre
    )
    cancha_id, score_cancha, candidatos_cancha = _mejor_coincidencia(
        fila["cancha_texto"], canchas, lambda item: item.nombre_cancha
    )
    cancha_sugerida = next((item for item in canchas if item.id == cancha_id), None)
    # Una cancha homónima en otra ciudad no debe prevalecer sobre la ciudad explícita del horario.
    if cancha_sugerida and ciudad_id is not None and cancha_sugerida.ciudad_id != ciudad_id:
        cancha_id, score_cancha = None, 0.0
    elif cancha_sugerida and ciudad_id is None and ciudad_texto:
        # Si el nombre de la ciudad no coincide con el catálogo, no inferirla desde
        # una cancha posiblemente mal registrada: requiere revisión humana.
        cancha_id, score_cancha = None, 0.0
    ciudad_origen_id, score_ciudad_origen, candidatos_origen = _mejor_coincidencia(
        fila.get("ciudad_origen_texto", ""), ciudades, lambda item: item.nombre
    )
    cancha = next((item for item in canchas if item.id == cancha_id), None)
    if ciudad_id is None and not ciudad_texto and cancha:
        ciudad_id = cancha.ciudad_id
        score_ciudad = 0.85
    arbitro = next((item for item in arbitros if item.id == arbitro_id), None)
    if ciudad_origen_id is None and arbitro and not fila.get("ciudad_origen_texto"):
        ciudad_origen_id = arbitro.ciudad_origen_id
        score_ciudad_origen = 1.0
    ciudad_origen = next((item for item in ciudades if item.id == ciudad_origen_id), None)
    ciudad_destino = next((item for item in ciudades if item.id == ciudad_id), None)
    scores = {
        "fecha": 1.0 if fila["fecha"] else 0.0,
        "hora": 1.0 if fila["hora"] else 0.0,
        "arbitro": score_arbitro,
        "cancha": score_cancha,
        "ciudad": score_ciudad,
        "ciudad_origen": score_ciudad_origen,
    }
    return {
        **fila,
        "arbitro": f"{arbitro.nombre} {arbitro.apellido}" if arbitro else fila.get("arbitro_texto", ""),
        "ciudad_origen": ciudad_origen.nombre if ciudad_origen else fila.get("ciudad_origen_texto", ""),
        "ciudad_destino": ciudad_destino.nombre if ciudad_destino else ciudad_texto,
        "cancha": cancha.nombre_cancha if cancha else fila.get("cancha_texto", ""),
        "arbitro_id": arbitro_id,
        "cancha_id": cancha_id,
        "ciudad_id": ciudad_id,
        "ciudad_destino_id": ciudad_id,
        "ciudad_origen_id": ciudad_origen_id,
        "candidatos_arbitro": candidatos_arbitro,
        "candidatos_cancha": candidatos_cancha,
        "candidatos_ciudad": candidatos_ciudad,
        "candidatos_ciudad_origen": candidatos_origen,
        "confianza_campos": scores,
        "confianza": round(sum(scores.values()) / len(scores), 3),
    }


def _validar_preview(filas: list[dict]) -> list[dict]:
    """Añade validaciones locales de fecha, hora, duplicados y confianza."""
    vistos = {}
    for indice, fila in enumerate(filas):
        errores = []
        fecha, hora = fila.get("fecha", ""), fila.get("hora", "")
        try:
            datetime.strptime(fecha, "%Y-%m-%d")
        except (TypeError, ValueError):
            errores.append("fecha ausente o inválida; usa YYYY-MM-DD")
        try:
            datetime.strptime(hora, "%H:%M")
        except (TypeError, ValueError):
            errores.append("hora ausente o inválida; usa HH:MM en formato 24 horas")

        arbitro_key = fila.get("arbitro_id") or _clave(fila.get("arbitro_texto", ""))
        clave_duplicado = (arbitro_key, fecha, hora)
        if arbitro_key and fecha and hora and clave_duplicado in vistos:
            errores.append(f"duplicado de la fila {vistos[clave_duplicado] + 1} para árbitro, fecha y hora")
            fila["duplicado"] = True
            filas[vistos[clave_duplicado]]["duplicado"] = True
        elif arbitro_key and fecha and hora:
            vistos[clave_duplicado] = indice

        try:
            confianza = float(fila.get("confianza", 0))
        except (TypeError, ValueError):
            confianza = 0.0
        fila["confianza_baja"] = confianza < 0.6
        fila["duplicado"] = fila.get("duplicado", False)
        fila["errores_validacion"] = errores
        fila["marcada_revision"] = bool(errores or fila["confianza_baja"])
    return filas


def _normalizar_filas_tabulares(frame) -> list[dict]:
    columnas = list(frame.columns)
    mapa = {tipo: _columna(columnas, tipo) for tipo in ALIAS}
    if not mapa["fecha"] and not mapa["fecha_hora"]:
        raise ValueError("El archivo debe incluir una columna de fecha o fecha/hora.")
    if not mapa["cancha"]:
        raise ValueError("El archivo debe incluir una columna de cancha, estadio o sede.")

    filas = []
    for indice, registro in frame.iterrows():
        fecha_valor = registro.get(mapa["fecha_hora"]) if mapa["fecha_hora"] else registro.get(mapa["fecha"])
        hora_valor = registro.get(mapa["hora"]) if mapa["hora"] else ""
        fecha, hora = _fecha_hora(fecha_valor, hora_valor)
        cantidad_dias = _numero_opcional(registro.get(mapa["cantidad_dias"])) if mapa["cantidad_dias"] else None
        cantidad_noches = _numero_opcional(registro.get(mapa["cantidad_noches"])) if mapa["cantidad_noches"] else None
        viaticos_total = _numero_opcional(registro.get(mapa["viaticos_total"])) if mapa["viaticos_total"] else None
        viatico_por_dia = (
            viaticos_total / cantidad_dias
            if viaticos_total is not None and cantidad_dias and cantidad_dias > 0 else None
        )
        fila = {
            "fila_origen": int(indice) + 2,
            "evidencia": f"Fila {int(indice) + 2}",
            "fecha": fecha,
            "hora": hora,
            "arbitro_texto": _texto_compacto(registro.get(mapa["arbitro"])) if mapa["arbitro"] else "",
            "cancha_texto": _texto_compacto(registro.get(mapa["cancha"])) if mapa["cancha"] else "",
            "ciudad_texto": _texto_compacto(registro.get(mapa["ciudad"])) if mapa["ciudad"] else "",
            "ciudad_origen_texto": _texto_compacto(registro.get(mapa["ciudad_origen"])) if mapa["ciudad_origen"] else "",
            "medio_transporte": _normalizar_medio(registro.get(mapa["tipo_traslado"])) if mapa["tipo_traslado"] else None,
        }
        if cantidad_dias is not None:
            fila["cantidad_dias"] = int(cantidad_dias) if cantidad_dias.is_integer() else cantidad_dias
        if cantidad_noches is not None:
            fila["cantidad_noches"] = int(cantidad_noches) if cantidad_noches.is_integer() else cantidad_noches
        if viatico_por_dia is not None:
            fila["viatico_por_dia"] = viatico_por_dia
        if mapa["hospedaje_total"]:
            hospedaje = _numero_opcional(registro.get(mapa["hospedaje_total"]))
            if hospedaje is not None:
                fila["hospedaje_total"] = hospedaje
        if not any(fila[campo] for campo in ("fecha", "hora", "arbitro_texto", "cancha_texto", "ciudad_texto")):
            continue
        filas.append(fila)
    return filas


def _fusionar_logistica_excel(filas_ia: list[dict], filas_fuente: list[dict]) -> list[dict]:
    """Usa valores numéricos explícitos del libro para evitar perderlos en la extracción LLM."""
    if not filas_ia or not filas_fuente:
        return filas_ia
    for fila in filas_ia:
        referencia = re.search(r"(?:\bR\s*|\bfila\s+)(\d+)\b", fila.get("evidencia", ""), flags=re.IGNORECASE)
        candidatas = []
        if referencia and not fila.get("evidencia", "").lower().startswith("fila de origen"):
            numero_fila = int(referencia.group(1))
            candidatas = [fuente for fuente in filas_fuente if fuente.get("fila_origen") == numero_fila]
        if not candidatas:
            candidatas = [
                fuente for fuente in filas_fuente
                if fuente.get("fecha") == fila.get("fecha") and fuente.get("hora") == fila.get("hora")
            ]
        if len(candidatas) > 1:
            nombre_cancha = _clave(fila.get("cancha_texto", ""))
            puntuadas = sorted(
                ((SequenceMatcher(None, nombre_cancha, _clave(c.get("cancha_texto", ""))).ratio(), c) for c in candidatas),
                key=lambda par: par[0], reverse=True,
            )
            if not nombre_cancha or puntuadas[0][0] < 0.5 or puntuadas[0][0] - puntuadas[1][0] < 0.05:
                continue
            fuente = puntuadas[0][1]
        elif candidatas:
            fuente = candidatas[0]
        else:
            continue

        for campo in ("cantidad_dias", "cantidad_noches", "viatico_por_dia", "hospedaje_total", "medio_transporte"):
            if fuente.get(campo) is not None:
                fila[campo] = fuente[campo]
        if fuente.get("ciudad_texto"):
            # Para libros tabulares, la celda CIUDAD explícita es más fiable que
            # una sugerencia geográfica alucinada por el modelo.
            fila["ciudad_texto"] = fuente["ciudad_texto"]
            fila["ciudad_destino_texto"] = fuente["ciudad_texto"]
        if fuente.get("ciudad_origen_texto"):
            fila["ciudad_origen_texto"] = fuente["ciudad_origen_texto"]
        if fuente.get("evidencia"):
            fila["evidencia"] = fuente["evidencia"]
    return filas_ia


def extraer_horario_arbitro(archivo_o_imagen, filename: str | None = None) -> dict:
    """Extrae filas sin escribir datos en la base; devuelve IDs sugeridos."""
    if hasattr(archivo_o_imagen, "read"):
        archivo = archivo_o_imagen
        filename = filename or getattr(archivo, "filename", "")
        contenido = archivo.read()
    else:
        contenido = archivo_o_imagen
    if not isinstance(contenido, (bytes, bytearray)):
        raise ValueError("El archivo debe recibirse como bytes o como archivo abierto.")
    extension = (filename or "").lower().rsplit(".", 1)[-1]
    extension = f".{extension}" if extension else ""
    advertencias, proveedor, modelo = [], "tabular", None
    arbitros = Arbitro.query.filter_by(activo=True).order_by(Arbitro.id).all()
    canchas = UbicacionCancha.query.order_by(UbicacionCancha.id).all()
    ciudades = Ciudad.query.order_by(Ciudad.id).all()
    catalogo = _catalogo_canonico(arbitros, ciudades, canchas)

    if extension == ".xlsx":
        error_ia = None
        if os.getenv("VISION_PROVIDER", "simulado").strip().lower() not in {"simulado", "emulado", "local"}:
            try:
                resultado_ia = extraer_filas_excel_texto(_texto_completo_excel(bytes(contenido)), catalogo=catalogo)
                origen = _filas_desde_ia(resultado_ia)
                proveedor, modelo = resultado_ia["proveedor"], resultado_ia.get("modelo")
                advertencias.extend(resultado_ia.get("advertencias", []))
                if not origen:
                    raise ValueError("La IA no identificó filas de partidos en el libro.")
            except ValueError as error:
                error_ia = error
                origen = []
        else:
            origen = []
        if not origen:
            try:
                origen = _normalizar_filas_excel(bytes(contenido))
                proveedor, modelo = "lectura local", None
                if error_ia:
                    advertencias.append(
                        f"No se pudo usar la IA ({error_ia}). Se aplicó la detección local por encabezados. "
                        "Revisa cuidadosamente los datos antes de confirmar."
                    )
            except ValueError as error_local:
                if error_ia:
                    raise ValueError(
                        f"La IA no pudo interpretar el libro ({error_ia}) y la lectura local tampoco encontró "
                        f"una tabla reconocible ({error_local}). Revisa el archivo o completa la plantilla manualmente."
                    ) from error_local
                raise
        if proveedor != "lectura local" and origen:
            try:
                filas_fuente = _normalizar_filas_excel(bytes(contenido))
                _fusionar_logistica_excel(origen, filas_fuente)
            except ValueError:
                # El modelo puede leer diseños no tabulares que no admite el parser local;
                # en ese caso se conserva la lectura IA y su revisión humana.
                pass
    elif extension == ".csv":
        error_ia = None
        if os.getenv("VISION_PROVIDER", "simulado").strip().lower() not in {"simulado", "emulado", "local"}:
            try:
                texto_csv = _texto_csv_para_ia(bytes(contenido))
                resultado_ia = extraer_filas_excel_texto(texto_csv, catalogo=catalogo)
                origen = _filas_desde_ia(resultado_ia)
                proveedor, modelo = resultado_ia["proveedor"], resultado_ia.get("modelo")
                advertencias.extend(resultado_ia.get("advertencias", []))
                if not origen:
                    raise ValueError("La IA no encontró partidos en el CSV.")
            except ValueError as error:
                error_ia = error
                origen = []
        else:
            origen = []
        if not origen:
            origen = _normalizar_filas_tabulares(_leer_tabular(bytes(contenido), extension))
            proveedor, modelo = "lectura local", None
            if error_ia:
                advertencias.append(
                    f"No se pudo usar la IA ({error_ia}). Se aplicó la lectura tabular local. "
                    "Revisa los campos antes de confirmar."
                )
    elif extension in {".png", ".jpg", ".jpeg"}:
        mime_type = "image/png" if extension == ".png" else "image/jpeg"
        resultado = extraer_filas_imagen(bytes(contenido), filename or "horario", mime_type, catalogo=catalogo)
        proveedor, modelo = resultado["proveedor"], resultado.get("modelo")
        advertencias.extend(resultado.get("advertencias", []))
        origen = _filas_desde_ia(resultado)
        for fila in origen:
            if fila["evidencia"].startswith("Fila de origen"):
                fila["evidencia"] = f"Imagen {filename or 'horario'} · fila visual {fila['fila_origen']}"
    else:
        raise ValueError("Formato no permitido. Adjunta un .xlsx, .csv, .png, .jpg o .jpeg.")

    if not origen:
        raise ValueError("No se encontraron filas de horario para previsualizar.")
    filas = []
    for item in origen:
        origen_confianza = item.pop("_confianza_origen", None)
        fila = _resolver_sugerencias(item, arbitros, canchas, ciudades)
        if origen_confianza is not None:
            try:
                confianza_origen = max(0.0, min(1.0, float(origen_confianza)))
            except (TypeError, ValueError):
                confianza_origen = 0.5
            fila["confianza"] = min(fila["confianza"], confianza_origen)
            fila["confianza_campos"] = {
                campo: min(score, confianza_origen)
                for campo, score in fila["confianza_campos"].items()
            }
        filas.append(fila)
    filas = _validar_preview(filas)
    advertencias_validacion = sum(bool(fila["errores_validacion"] or fila["confianza_baja"]) for fila in filas)
    if advertencias_validacion:
        advertencias.append(
            f"{advertencias_validacion} fila(s) requieren revisión por fechas/horas inválidas, duplicados o confianza baja."
        )
    return {
        "solo_lectura": True,
        "origen": proveedor,
        "modelo": modelo,
        "advertencias": advertencias,
        "filas": filas,
    }
